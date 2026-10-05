from pptx import Presentation

prs = Presentation()

# Title Slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "ResumeAI"
subtitle.text = "AI-Powered ATS Resume Analyzer & Live Editor\nMini Project Presentation"

# Slide 2: Problem Statement
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Problem Statement"
tf = body.text_frame
tf.text = "Job seekers struggle to pass ATS (Applicant Tracking Systems)."
p = tf.add_paragraph()
p.text = "Candidates lack feedback on missing keywords or formatting."
p = tf.add_paragraph()
p.text = "Generic resumes fail to match specific job descriptions."
p = tf.add_paragraph()
p.text = "Editing resumes iteratively is tedious and time-consuming."

# Slide 3: Solution
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Our Solution: ResumeAI"
tf = body.text_frame
tf.text = "An intelligent web application that:"
p = tf.add_paragraph()
p.text = "Extracts and parses text from uploaded PDF resumes."
p = tf.add_paragraph()
p.text = "Compares candidate skills against target roles or job descriptions."
p = tf.add_paragraph()
p.text = "Provides actionable feedback, ATS scoring, and keyword matching."
p = tf.add_paragraph()
p.text = "Includes a Live Editor and 120+ dynamically generated templates."

# Slide 4: Key Features
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Key Features"
tf = body.text_frame
tf.text = "1. NLP Analysis: Keyword extraction and scoring."
p = tf.add_paragraph()
p.text = "2. Job Matching: Compares resume against specific role profiles."
p = tf.add_paragraph()
p.text = "3. Live Resume Editor: Edit parsed text directly in the browser."
p = tf.add_paragraph()
p.text = "4. Template Engine: Generates 120+ unique visual resumes."
p = tf.add_paragraph()
p.text = "5. Progress Tracking: Compare 'Before' and 'After' scores."

# Slide 5: Technology Stack
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Technology Stack"
tf = body.text_frame
tf.text = "Frontend:"
p = tf.add_paragraph()
p.text = "- HTML5, CSS3, Vanilla JavaScript (No heavy frameworks)"
p = tf.add_paragraph()
p.text = "Backend:"
p = tf.add_paragraph()
p.text = "- Python & Flask"
p = tf.add_paragraph()
p.text = "- PyPDF2 for text extraction"
p = tf.add_paragraph()
p.text = "Deployment:"
p = tf.add_paragraph()
p.text = "- Stateless Vercel Serverless Functions"

# Slide 6: Architecture
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Architecture & Workflow"
tf = body.text_frame
tf.text = "1. User uploads PDF -> Flask backend extracts raw text."
p = tf.add_paragraph()
p.text = "2. Backend calculates keyword overlap and section completeness."
p = tf.add_paragraph()
p.text = "3. Frontend displays dashboard with scores and weaknesses."
p = tf.add_paragraph()
p.text = "4. User edits text in Live Editor -> hits stateless /analyze endpoint."
p = tf.add_paragraph()
p.text = "5. User selects template -> dynamically populates HTML -> saves as PDF."

# Slide 7: Conclusion
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Conclusion & Future Scope"
tf = body.text_frame
tf.text = "ResumeAI bridges the gap between candidates and recruiters."
p = tf.add_paragraph()
p.text = "Future enhancements:"
p = tf.add_paragraph()
p.text = "- Deep Learning for semantic matching (instead of keyword matching)."
p = tf.add_paragraph()
p.text = "- OAuth login to save multiple resumes."
p = tf.add_paragraph()
p.text = "- Cover letter generation based on analysis."

prs.save("ResumeAI_Presentation.pptx")
print("Presentation generated!")
